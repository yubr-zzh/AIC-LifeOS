import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def ensure_style(document, name, font_name="Arial", font_size=10, color=None):
    styles = document.styles
    if name in styles:
        style = styles[name]
    else:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font_name
    style.font.size = Pt(font_size)
    if color:
        style.font.color.rgb = RGBColor(*color)
    return style


def add_inline_runs(paragraph, text):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_table(document, rows):
    cells = [split_table_row(row) for row in rows]
    if len(cells) < 2:
        return
    if is_separator_row(cells[1]):
        body = [cells[0], *cells[2:]]
    else:
        body = cells
    table = document.add_table(rows=0, cols=max(len(row) for row in body))
    table.style = "Table Grid"
    for row_index, row_cells in enumerate(body):
        row = table.add_row()
        for col_index, value in enumerate(row_cells):
            cell = row.cells[col_index]
            paragraph = cell.paragraphs[0]
            add_inline_runs(paragraph, value.strip())
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def split_table_row(row):
    stripped = row.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def convert(md_path, docx_path):
    text = Path(md_path).read_text(encoding="utf-8")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10.5)
    ensure_style(document, "Code Block", "Consolas", 8.5, color=(40, 40, 40))

    in_code = False
    code_lines = []
    table_lines = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        paragraph = document.add_paragraph(style="Code Block")
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.add_run("\n".join(code_lines))
        code_lines = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            add_table(document, table_lines)
            table_lines = []

    for raw_line in text.splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines.append(line)
            continue
        flush_table()

        if not line.strip():
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 4)
            paragraph = document.add_heading(heading.group(2), level=level)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue

        bullet = re.match(r"^\s*-\s+(.*)$", line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet.group(1))
            continue

        number = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if number:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, number.group(1))
            continue

        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, line)

    flush_code()
    flush_table()
    document.save(docx_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: md_to_docx.py input.md output.docx")
    convert(sys.argv[1], sys.argv[2])

